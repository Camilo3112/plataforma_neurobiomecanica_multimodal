"""
╔══════════════════════════════════════════════════════════════════════════════╗
║             19 ZONAS ANATÓMICAS · CORRELACIÓN LOCAL · EXCEL/NIfTI            ║
╚══════════════════════════════════════════════════════════════════════════════╝

Archivo: src/zonas_correlacion_runner.py
Versión: v3.21.23

Descripción
-----------
Módulo de análisis longitudinal para delimitar 19 regiones anatómicas sobre el
T1/rAnatomico del paciente y calcular similitud local Antes vs Después. Genera
máscaras NIfTI, mapas de baja correlación, labelmaps para ITK-SNAP/3D Slicer y
libros Excel con volúmenes y métricas de correlación por región.

Fundamento físico-matemático implementado
-----------------------------------------
1. Representación espacial de neuroimagen
   Cada volumen NIfTI/MGZ se interpreta mediante una matriz afín homogénea A de
   tamaño 4x4. La relación entre coordenadas discretas del voxel i=(i,j,k,1)^T
   y coordenadas físicas RAS-mm se expresa como:

       x_RAS = A · i

   Las etiquetas anatómicas se remuestrean al espacio del T1 usando interpolación
   por vecino más cercano, porque los labels son variables categóricas discretas
   y no intensidades continuas.

2. Delimitación anatómica de 19 zonas
   Las regiones se obtienen desde FreeSurfer aparc+aseg/aseg. Algunas zonas son
   etiquetas directas del atlas Desikan-Killiany; otras son aproximaciones
   geométricas por cuantiles espaciales o bandas mediales/laterales. Para una
   máscara M, un subconjunto anterior, posterior, superior o inferior se define
   con cuantiles de las coordenadas físicas:

       M_q = { v ∈ M : coord_axis(v) ≥ Q_{1-q} }  para porciones superiores/anterior
       M_q = { v ∈ M : coord_axis(v) ≤ Q_q }      para porciones inferiores/posterior

   Los volúmenes se calculan por integración discreta:

       V_mm3 = N_voxeles · Δx · Δy · Δz
       V_ml  = V_mm3 / 1000

3. Registro longitudinal rígido Después -> Antes
   El volumen Después se registra sobre Antes con una transformación rígida 3D:

       T(x) = R · x + t

   donde R es una matriz de rotación 3D y t un vector de traslación. La métrica
de optimización usa información mutua de Mattes, útil cuando existen cambios de
   intensidad entre adquisiciones.

4. Normalización robusta de intensidades
   Antes de comparar intensidades se usa z-score robusto dentro de la unión de
   ROIs, basado en mediana y MAD:

       z(v) = (I(v) - mediana(I_M)) / (1.4826 · MAD(I_M))

   Esto reduce sensibilidad a outliers y diferencias globales de escala.

5. Correlación local 2D por cortes
   Para cada voxel se calcula una correlación local tipo Pearson/NCC dentro de
   ventanas 2D de tamaño w×w recorriendo todos los cortes de los tres ejes. La
   correlación local se calcula como:

       r = cov(A,B) / sqrt(var(A) · var(B))

   Los mapas de los tres planos se combinan promediando solo valores válidos:

       r_3planos(v) = promedio{ r_eje0(v), r_eje1(v), r_eje2(v) }

   Una baja correlación indica menor similitud local después del registro y debe
   interpretarse junto con control visual sobre el T1 de referencia.

Salidas principales
-------------------
- regiones_adicionales_T1/individuales/*.nii
- regiones_adicionales_multietiqueta_T1.nii
- volumen_regiones_adicionales_por_hemisferio.csv
- regiones_adicionales_reporte.json
- correlacion_antes_despues_19_zonas/correlacion_local_3planos_T1.nii.gz
- correlacion_antes_despues_19_zonas/baja_correlacion_score_T1.nii.gz
- correlacion_antes_despues_19_zonas/baja_correlacion_zonas_labelmap_T1.nii.gz
- resultados/resumen_volumenes_19_zonas.xlsx
- resultados/resumen_correlacion_19_zonas.xlsx
"""

###############################################################################################################################
###############################################################################################################################

# ══════════════════════════════════════════════════════════════════════════════
#  1. IMPORTACIONES Y CONFIGURACIÓN GENERAL
# ══════════════════════════════════════════════════════════════════════════════

from __future__ import annotations

import csv
import json
import math
import os
import shutil
import subprocess
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, Optional, Tuple

import nibabel as nib
import numpy as np
import pandas as pd
import SimpleITK as sitk
from nibabel.processing import resample_from_to
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from scipy.ndimage import uniform_filter
from scipy.stats import pearsonr

from .config import PipelineConfig
from .paths import resolve_stage_dir, stage_key

WINDOW_SIZE = int(os.environ.get("VCE_ZONAS_WINDOW_SIZE", "7"))
LOW_CORR_THRESHOLD = float(os.environ.get("VCE_ZONAS_LOW_CORR_THRESHOLD", "0.50"))


###############################################################################################################################
###############################################################################################################################

# ══════════════════════════════════════════════════════════════════════════════
#  2. DEFINICIÓN DE ATLAS, ETIQUETAS Y REGIONES
# ══════════════════════════════════════════════════════════════════════════════

DK = {
    "lh_caudalmiddlefrontal": 1003, "rh_caudalmiddlefrontal": 2003,
    "lh_fusiform": 1007, "rh_fusiform": 2007,
    "lh_inferiorparietal": 1008, "rh_inferiorparietal": 2008,
    "lh_inferiortemporal": 1009, "rh_inferiortemporal": 2009,
    "lh_lateraloccipital": 1011, "rh_lateraloccipital": 2011,
    "lh_lateralorbitofrontal": 1012, "rh_lateralorbitofrontal": 2012,
    "lh_lingual": 1013, "rh_lingual": 2013,
    "lh_medialorbitofrontal": 1014, "rh_medialorbitofrontal": 2014,
    "lh_paracentral": 1017, "rh_paracentral": 2017,
    "lh_parsopercularis": 1018, "rh_parsopercularis": 2018,
    "lh_parsorbitalis": 1019, "rh_parsorbitalis": 2019,
    "lh_parstriangularis": 1020, "rh_parstriangularis": 2020,
    "lh_pericalcarine": 1021, "rh_pericalcarine": 2021,
    "lh_postcentral": 1022, "rh_postcentral": 2022,
    "lh_precentral": 1024, "rh_precentral": 2024,
    "lh_precuneus": 1025, "rh_precuneus": 2025,
    "lh_rostralmiddlefrontal": 1027, "rh_rostralmiddlefrontal": 2027,
    "lh_superiorfrontal": 1028, "rh_superiorfrontal": 2028,
    "lh_superiorparietal": 1029, "rh_superiorparietal": 2029,
    "lh_supramarginal": 1031, "rh_supramarginal": 2031,
    "lh_frontalpole": 1032, "rh_frontalpole": 2032,
}

ASEG = {
    "left_cerebellum_wm": 7,
    "left_cerebellum_cortex": 8,
    "right_cerebellum_wm": 46,
    "right_cerebellum_cortex": 47,
}

REGION_ORDER = [
    ("corteza_prefrontal", 1, (240, 80, 80)),
    ("corteza_prefrontal_dorsolateral", 2, (255, 140, 70)),
    ("vermis_cerebeloso", 3, (180, 90, 220)),
    ("corteza_motora_primaria", 4, (250, 30, 30)),
    ("corteza_somatosensitiva_primaria", 5, (250, 200, 40)),
    ("area_motora_suplementaria", 6, (255, 120, 120)),
    ("precuneo", 7, (70, 170, 220)),
    ("corteza_intracalcarina", 8, (50, 120, 250)),
    ("giro_lingual", 9, (140, 70, 250)),
    ("giro_fusiforme_occipital", 10, (255, 50, 180)),
    ("corteza_occipital_lateral_superior", 11, (100, 200, 255)),
    ("corteza_occipital_lateral_inferior", 12, (60, 120, 220)),
    ("giro_angular", 13, (120, 220, 120)),
    ("lobulo_parietal_superior", 14, (70, 240, 160)),
    ("area_somatosensorial_secundaria", 15, (255, 170, 70)),
    ("cerebelo_lobulo_vi", 16, (160, 60, 240)),
    ("area_motora_presuplementaria", 17, (255, 150, 150)),
    ("area_somatosensorial_secundaria_posterior", 18, (220, 200, 100)),
    ("giro_temporal_inferior_posterior", 19, (200, 100, 220)),
]

MIDLINE_REGIONS = {"vermis_cerebeloso"}


@dataclass
class RegionResult:
    """Máscara, color y trazabilidad de una región anatómica."""

    name: str
    label_id: int
    color: Tuple[int, int, int]
    mask: np.ndarray
    source: str
    note: str


###############################################################################################################################
###############################################################################################################################

# ══════════════════════════════════════════════════════════════════════════════
#  3. UTILIDADES DE RUTAS, LOGS Y NIfTI
# ══════════════════════════════════════════════════════════════════════════════


def log(msg: str) -> None:
    print(msg, flush=True)


def stamp() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def sanitize_name(name: str) -> str:
    return (
        name.replace("á", "a").replace("é", "e").replace("í", "i")
        .replace("ó", "o").replace("ú", "u").replace("ñ", "n")
    )


def patient_number(patient: str) -> str:
    return "".join(ch for ch in str(patient) if ch.isdigit()) or str(patient).replace(" ", "_")


def subject_id_for(patient: str, stage: str) -> str:
    return f"{patient.replace(' ', '')}_{stage}"


def first_existing(paths: Iterable[Path]) -> Optional[Path]:
    for p in paths:
        if p and Path(p).exists():
            return Path(p)
    return None


def t1_candidates(stage_dir: Optional[Path]) -> list[Path]:
    if stage_dir is None:
        return []
    stage_dir = Path(stage_dir)
    explicit = [
        stage_dir / "ResultadosFuncional" / "REFORMATEO" / "rAnatomico.nii",
        stage_dir / "Resultados funcional" / "REFORMATEO" / "rAnatomico.nii",
        stage_dir / "ResultadosFuncional" / "Resultados" / "Resultados" / "rAnatomico.nii",
        stage_dir / "Resultados funcional" / "Resultados" / "Resultados" / "rAnatomico.nii",
        stage_dir / "ResultadosFuncional" / "REFORMATEO" / "rT1.nii",
        stage_dir / "Resultados funcional" / "REFORMATEO" / "rT1.nii",
        stage_dir / "RESONANCIA" / "REFORMATEO" / "rAnatomico.nii",
        stage_dir / "RESONANCIA" / "REFORMATEO" / "rT1.nii",
        stage_dir / "Resonancia" / "REFORMATEO" / "rAnatomico.nii",
        stage_dir / "Resonancia" / "REFORMATEO" / "rT1.nii",
    ]
    names = {"ranatomico.nii", "ranatomico.nii.gz", "rt1.nii", "rt1.nii.gz", "t1.nii", "t1.nii.gz"}
    recursive: list[Path] = []
    for root in [stage_dir / "ResultadosFuncional", stage_dir / "Resultados funcional", stage_dir / "RESONANCIA", stage_dir / "Resonancia", stage_dir]:
        if not root.exists():
            continue
        try:
            for f in root.rglob("*.nii*"):
                lname = f.name.lower()
                if lname in names and not any(bad in lname for bad in ["mask", "roi", "dwi", "fa", "adc", "bold", "func"]):
                    recursive.append(f)
        except Exception:
            continue
    ordered: list[Path] = []
    seen: set[str] = set()
    for f in explicit + recursive:
        key = str(f.resolve()) if f.exists() else str(f)
        if key not in seen:
            seen.add(key)
            ordered.append(f)
    return ordered


def load_nifti(path: Path) -> tuple[nib.Nifti1Image, np.ndarray]:
    img = nib.load(str(path))
    return img, np.asarray(img.dataobj)


def voxel_volume_mm3(img: nib.Nifti1Image) -> float:
    zooms = img.header.get_zooms()[:3]
    return float(zooms[0] * zooms[1] * zooms[2])


def save_nifti(reference_img: nib.Nifti1Image, data: np.ndarray, path: Path, dtype=np.float32) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    header = reference_img.header.copy()
    header.set_data_dtype(dtype)
    out = nib.Nifti1Image(data.astype(dtype), reference_img.affine, header)
    out.set_qform(reference_img.affine, code=1)
    out.set_sform(reference_img.affine, code=1)
    nib.save(out, str(path))


def geometry_matches(a: nib.Nifti1Image, b: nib.Nifti1Image, atol: float = 1e-3) -> bool:
    return tuple(a.shape[:3]) == tuple(b.shape[:3]) and np.allclose(a.affine, b.affine, atol=atol)


###############################################################################################################################
###############################################################################################################################

# ══════════════════════════════════════════════════════════════════════════════
#  4. EXPORTACIÓN FREESURFER -> T1 NATIVO
# ══════════════════════════════════════════════════════════════════════════════


def run_realtime(cmd: list[str], log_path: Path, env: Optional[dict[str, str]] = None) -> int:
    log_path.parent.mkdir(parents=True, exist_ok=True)
    log("\n" + "─" * 92)
    log("COMANDO: " + " ".join(str(x) for x in cmd))
    log("LOG: " + str(log_path))
    log("─" * 92)
    with log_path.open("w", encoding="utf-8", errors="replace") as fh:
        fh.write("COMANDO: " + " ".join(str(x) for x in cmd) + "\n")
        proc = subprocess.Popen(
            [str(x) for x in cmd],
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            env=env or os.environ.copy(),
            bufsize=1,
            universal_newlines=True,
        )
        assert proc.stdout is not None
        for line in proc.stdout:
            print(line, end="", flush=True)
            fh.write(line)
            fh.flush()
        proc.wait()
        fh.write(f"\nCODIGO_SALIDA={proc.returncode}\n")
    log(f"Subproceso terminado con código {proc.returncode}")
    return int(proc.returncode or 0)


def resample_label_to_t1_with_nibabel(source_label: Path, t1_path: Path, output_path: Path) -> Path:
    """Remuestreo categórico por vecino más cercano usando las afines del archivo."""
    t1_img = nib.load(str(t1_path))
    source_img = nib.load(str(source_label))
    resampled = resample_from_to(source_img, (t1_img.shape[:3], t1_img.affine), order=0)
    data = np.asanyarray(resampled.dataobj).astype(np.int32)
    save_nifti(t1_img, data, output_path, dtype=np.int32)
    return output_path


def export_label_with_freesurfer_or_nibabel(source_label: Path, t1_path: Path, output_path: Path, log_path: Path) -> Path:
    """Exporta aparc+aseg/aseg al espacio T1.

    Primero intenta `mri_vol2vol --regheader --nearest`, porque conserva la
    relación geométrica esperada en sujetos generados por recon-all. Si el
    comando no está disponible o falla, se usa remuestreo NIfTI/MGZ con nibabel.
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)
    mri_vol2vol = shutil.which("mri_vol2vol", path=os.environ.get("PATH"))
    if mri_vol2vol:
        cmd = [
            mri_vol2vol,
            "--mov", str(source_label),
            "--targ", str(t1_path),
            "--regheader",
            "--o", str(output_path),
            "--nearest",
        ]
        code = run_realtime(cmd, log_path)
        if code == 0 and output_path.exists():
            return output_path
    log("[AVISO] Se usará remuestreo nibabel por vecino más cercano.")
    return resample_label_to_t1_with_nibabel(source_label, t1_path, output_path)


def ensure_freesurfer_labels_in_t1(
    *,
    cfg: PipelineConfig,
    patient: str,
    stage: str,
    t1_path: Path,
    output_dir: Path,
) -> tuple[Path, Path]:
    """Garantiza aparc_aseg_T1.nii.gz y aseg_T1.nii.gz para generar las 19 zonas."""
    source_dir = output_dir / "fuentes"
    source_dir.mkdir(parents=True, exist_ok=True)
    aparc_t1 = source_dir / "aparc_aseg_T1.nii.gz"
    aseg_t1 = source_dir / "aseg_T1.nii.gz"

    if aparc_t1.exists() and aseg_t1.exists() and not cfg.force:
        try:
            t1_img = nib.load(str(t1_path))
            if geometry_matches(t1_img, nib.load(str(aparc_t1))) and geometry_matches(t1_img, nib.load(str(aseg_t1))):
                return aparc_t1, aseg_t1
        except Exception:
            pass

    subjects_dir = Path(os.environ.get("VCE_FREESURFER_SUBJECTS_DIR", str(Path.home() / "freesurfer_subjects")))
    sid = subject_id_for(patient, stage)
    fs_mri = subjects_dir / sid / "mri"
    aparc_mgz = fs_mri / "aparc+aseg.mgz"
    aseg_mgz = fs_mri / "aseg.mgz"

    if not aparc_mgz.exists() or not aseg_mgz.exists():
        raise FileNotFoundError(
            "No se encontró aparc+aseg.mgz/aseg.mgz para generar las 19 zonas. "
            f"Ejecuta primero cst_tronco o fs_mni_motor para {patient}/{stage}.\n"
            f"Esperado: {aparc_mgz}\nEsperado: {aseg_mgz}"
        )

    log(f"[{stamp()}] Exportando aparc+aseg y aseg al espacio T1: {patient} / {stage}")
    export_label_with_freesurfer_or_nibabel(aparc_mgz, t1_path, aparc_t1, output_dir / "logs" / "export_aparc_aseg_T1.log")
    export_label_with_freesurfer_or_nibabel(aseg_mgz, t1_path, aseg_t1, output_dir / "logs" / "export_aseg_T1.log")

    t1_img = nib.load(str(t1_path))
    if not geometry_matches(t1_img, nib.load(str(aparc_t1))):
        raise RuntimeError("aparc_aseg_T1.nii.gz no coincide geométricamente con el T1/rAnatomico.")
    if not geometry_matches(t1_img, nib.load(str(aseg_t1))):
        raise RuntimeError("aseg_T1.nii.gz no coincide geométricamente con el T1/rAnatomico.")
    return aparc_t1, aseg_t1


###############################################################################################################################
###############################################################################################################################

# ══════════════════════════════════════════════════════════════════════════════
#  5. OPERADORES GEOMÉTRICOS PARA CONSTRUIR REGIONES
# ══════════════════════════════════════════════════════════════════════════════


def mask_from_labels(seg: np.ndarray, labels: Iterable[int]) -> np.ndarray:
    return np.isin(seg, list(labels))


def world_coords(mask: np.ndarray, affine: np.ndarray) -> np.ndarray:
    ijk = np.argwhere(mask)
    if len(ijk) == 0:
        return np.empty((0, 3), dtype=np.float32)
    return nib.affines.apply_affine(affine, ijk)


def union_masks(*masks: np.ndarray) -> np.ndarray:
    result = np.zeros(masks[0].shape, dtype=bool)
    for m in masks:
        result |= m.astype(bool)
    return result


def estimate_midline_x(mask: np.ndarray, affine: np.ndarray) -> float:
    coords = world_coords(mask, affine)
    if len(coords) == 0:
        return 0.0
    x = coords[:, 0]
    return float((np.min(x) + np.max(x)) / 2.0)


def axis_quantile_mask(mask: np.ndarray, affine: np.ndarray, axis: int, keep: str, q: float) -> np.ndarray:
    coords = world_coords(mask, affine)
    if len(coords) == 0:
        return np.zeros(mask.shape, dtype=bool)
    vals = coords[:, axis]
    if keep in {"anterior", "superior"}:
        thr = np.quantile(vals, 1.0 - q)
        keep_idx = vals >= thr
    elif keep in {"posterior", "inferior"}:
        thr = np.quantile(vals, q)
        keep_idx = vals <= thr
    else:
        raise ValueError(keep)
    ijk = np.argwhere(mask)
    out = np.zeros(mask.shape, dtype=bool)
    sel = ijk[keep_idx]
    if len(sel):
        out[sel[:, 0], sel[:, 1], sel[:, 2]] = True
    return out


def hemispheric_split(mask: np.ndarray, affine: np.ndarray, side: str, midline_x: float, tol: float = 2.0) -> np.ndarray:
    coords = world_coords(mask, affine)
    if len(coords) == 0:
        return np.zeros(mask.shape, dtype=bool)
    x = coords[:, 0]
    ijk = np.argwhere(mask)
    keep = x <= midline_x + tol if side == "left" else x >= midline_x - tol
    out = np.zeros(mask.shape, dtype=bool)
    sel = ijk[keep]
    if len(sel):
        out[sel[:, 0], sel[:, 1], sel[:, 2]] = True
    return out


def medial_fraction(mask: np.ndarray, affine: np.ndarray, side: str, midline_x: float, fraction: float) -> np.ndarray:
    coords = world_coords(mask, affine)
    if len(coords) == 0:
        return np.zeros(mask.shape, dtype=bool)
    ijk = np.argwhere(mask)
    dist = np.abs(coords[:, 0] - midline_x)
    thr = np.quantile(dist, fraction)
    keep = dist <= thr
    x = coords[:, 0]
    keep &= x <= midline_x + 2.0 if side == "left" else x >= midline_x - 2.0
    out = np.zeros(mask.shape, dtype=bool)
    sel = ijk[keep]
    if len(sel):
        out[sel[:, 0], sel[:, 1], sel[:, 2]] = True
    return out


def lateral_fraction(mask: np.ndarray, affine: np.ndarray, side: str, midline_x: float, fraction: float) -> np.ndarray:
    coords = world_coords(mask, affine)
    if len(coords) == 0:
        return np.zeros(mask.shape, dtype=bool)
    ijk = np.argwhere(mask)
    dist = np.abs(coords[:, 0] - midline_x)
    thr = np.quantile(dist, 1.0 - fraction)
    keep = dist >= thr
    x = coords[:, 0]
    keep &= x <= midline_x + 2.0 if side == "left" else x >= midline_x - 2.0
    out = np.zeros(mask.shape, dtype=bool)
    sel = ijk[keep]
    if len(sel):
        out[sel[:, 0], sel[:, 1], sel[:, 2]] = True
    return out


def by_level_medial_band(mask: np.ndarray, affine: np.ndarray, midline_x: float, half_width_mm: float) -> np.ndarray:
    coords = world_coords(mask, affine)
    if len(coords) == 0:
        return np.zeros(mask.shape, dtype=bool)
    ijk = np.argwhere(mask)
    keep = np.abs(coords[:, 0] - midline_x) <= half_width_mm
    out = np.zeros(mask.shape, dtype=bool)
    sel = ijk[keep]
    if len(sel):
        out[sel[:, 0], sel[:, 1], sel[:, 2]] = True
    return out


def write_itksnap_labels(path: Path, regions: Iterable[RegionResult] = ()) -> None:
    lines = [
        "################################################",
        "# ITK-SnAP Label Description File",
        "# IDX   -R-  -G-  -B-  -A--  VIS MSH  LABEL",
        '0       0    0    0    0     0   0    "Clear Label"',
    ]
    if not regions:
        for name, label_id, color in REGION_ORDER:
            regions = [RegionResult(name, label_id, color, np.zeros((1, 1, 1), dtype=bool), "", "") for name, label_id, color in REGION_ORDER]
            break
    for rg in regions:
        r, g, b = rg.color
        lines.append(f'{rg.label_id:<7d} {r:<4d} {g:<4d} {b:<4d} 1     1   1    "{sanitize_name(rg.name)}"')
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


###############################################################################################################################
###############################################################################################################################

# ══════════════════════════════════════════════════════════════════════════════
#  6. CONSTRUCCIÓN DE LAS 19 REGIONES EN ESPACIO T1
# ══════════════════════════════════════════════════════════════════════════════


def build_regions(aparc_seg: np.ndarray, aseg_seg: np.ndarray, ref_img: nib.Nifti1Image) -> list[RegionResult]:
    affine = ref_img.affine
    brain_mask = aseg_seg > 0
    midline_x = estimate_midline_x(brain_mask, affine)

    precentral = mask_from_labels(aparc_seg, [DK["lh_precentral"], DK["rh_precentral"]])
    postcentral = mask_from_labels(aparc_seg, [DK["lh_postcentral"], DK["rh_postcentral"]])
    superiorfrontal = mask_from_labels(aparc_seg, [DK["lh_superiorfrontal"], DK["rh_superiorfrontal"]])
    rostralmf = mask_from_labels(aparc_seg, [DK["lh_rostralmiddlefrontal"], DK["rh_rostralmiddlefrontal"]])
    caudalmf = mask_from_labels(aparc_seg, [DK["lh_caudalmiddlefrontal"], DK["rh_caudalmiddlefrontal"]])
    frontalpole = mask_from_labels(aparc_seg, [DK["lh_frontalpole"], DK["rh_frontalpole"]])
    parsop = mask_from_labels(aparc_seg, [DK["lh_parsopercularis"], DK["rh_parsopercularis"]])
    parsorb = mask_from_labels(aparc_seg, [DK["lh_parsorbitalis"], DK["rh_parsorbitalis"]])
    parstri = mask_from_labels(aparc_seg, [DK["lh_parstriangularis"], DK["rh_parstriangularis"]])
    latof = mask_from_labels(aparc_seg, [DK["lh_lateralorbitofrontal"], DK["rh_lateralorbitofrontal"]])
    medof = mask_from_labels(aparc_seg, [DK["lh_medialorbitofrontal"], DK["rh_medialorbitofrontal"]])
    precuneus = mask_from_labels(aparc_seg, [DK["lh_precuneus"], DK["rh_precuneus"]])
    pericalcarine = mask_from_labels(aparc_seg, [DK["lh_pericalcarine"], DK["rh_pericalcarine"]])
    lingual = mask_from_labels(aparc_seg, [DK["lh_lingual"], DK["rh_lingual"]])
    fusiform = mask_from_labels(aparc_seg, [DK["lh_fusiform"], DK["rh_fusiform"]])
    latocc = mask_from_labels(aparc_seg, [DK["lh_lateraloccipital"], DK["rh_lateraloccipital"]])
    infpar = mask_from_labels(aparc_seg, [DK["lh_inferiorparietal"], DK["rh_inferiorparietal"]])
    suppar = mask_from_labels(aparc_seg, [DK["lh_superiorparietal"], DK["rh_superiorparietal"]])
    supramarginal = mask_from_labels(aparc_seg, [DK["lh_supramarginal"], DK["rh_supramarginal"]])
    inf_temp = mask_from_labels(aparc_seg, [DK["lh_inferiortemporal"], DK["rh_inferiortemporal"]])
    paracentral = mask_from_labels(aparc_seg, [DK["lh_paracentral"], DK["rh_paracentral"]])

    cerebellum = mask_from_labels(aseg_seg, [
        ASEG["left_cerebellum_wm"], ASEG["left_cerebellum_cortex"],
        ASEG["right_cerebellum_wm"], ASEG["right_cerebellum_cortex"],
    ])
    cerebellum_cortex = mask_from_labels(aseg_seg, [ASEG["left_cerebellum_cortex"], ASEG["right_cerebellum_cortex"]])

    regions: list[RegionResult] = []
    cmap = {name: (rid, col) for name, rid, col in REGION_ORDER}

    def add(name: str, mask: np.ndarray, source: str, note: str) -> None:
        rid, col = cmap[name]
        regions.append(RegionResult(name, rid, col, mask.astype(bool), source, note))

    sf_anterior = axis_quantile_mask(superiorfrontal, affine, 1, "anterior", 0.65)
    pfc = union_masks(rostralmf, caudalmf, frontalpole, parsop, parsorb, parstri, latof, medof, sf_anterior)
    add("corteza_prefrontal", pfc, "aparc+aseg", "Unión de regiones frontales no motoras; superior frontal restringido a porción anterior.")

    left_sf_lat = lateral_fraction(hemispheric_split(superiorfrontal, affine, "left", midline_x), affine, "left", midline_x, 0.55)
    right_sf_lat = lateral_fraction(hemispheric_split(superiorfrontal, affine, "right", midline_x), affine, "right", midline_x, 0.55)
    dlpfc = union_masks(rostralmf, caudalmf, left_sf_lat, right_sf_lat)
    add("corteza_prefrontal_dorsolateral", dlpfc, "aparc+aseg + división geométrica", "Middle frontal + fracción lateral del superior frontal.")

    vermis = by_level_medial_band(cerebellum, affine, midline_x, 7.0)
    add("vermis_cerebeloso", vermis, "aseg + banda medial", "Aproximación medial del cerebelo centrada en línea media.")

    add("corteza_motora_primaria", precentral, "aparc+aseg", "Giro precentral bilateral.")
    add("corteza_somatosensitiva_primaria", postcentral, "aparc+aseg", "Giro postcentral bilateral.")

    sma_source = union_masks(superiorfrontal, paracentral)
    sma_medial = union_masks(
        medial_fraction(hemispheric_split(sma_source, affine, "left", midline_x), affine, "left", midline_x, 0.65),
        medial_fraction(hemispheric_split(sma_source, affine, "right", midline_x), affine, "right", midline_x, 0.65),
    )
    sma = axis_quantile_mask(sma_medial, affine, 1, "posterior", 0.55)
    add("area_motora_suplementaria", sma, "aparc+aseg + cuantiles", "Porción medial posterior del complejo superior frontal/paracentral.")

    add("precuneo", precuneus, "aparc+aseg", "Precuneus bilateral.")
    add("corteza_intracalcarina", pericalcarine, "aparc+aseg", "Aproximación mediante pericalcarine.")
    add("giro_lingual", lingual, "aparc+aseg", "Giro lingual bilateral.")

    occ_fus = axis_quantile_mask(fusiform, affine, 1, "posterior", 0.55)
    add("giro_fusiforme_occipital", occ_fus, "aparc+aseg + cuantiles", "Mitad posterior del fusiforme.")

    add("corteza_occipital_lateral_superior", axis_quantile_mask(latocc, affine, 2, "superior", 0.50), "aparc+aseg + cuantiles", "Mitad superior del lateral occipital.")
    add("corteza_occipital_lateral_inferior", axis_quantile_mask(latocc, affine, 2, "inferior", 0.50), "aparc+aseg + cuantiles", "Mitad inferior del lateral occipital.")

    add("giro_angular", axis_quantile_mask(infpar, affine, 1, "posterior", 0.50), "aparc+aseg + cuantiles", "Mitad posterior del inferior parietal.")
    add("lobulo_parietal_superior", suppar, "aparc+aseg", "Superior parietal bilateral.")

    s2_source = union_masks(postcentral, supramarginal)
    s2_inferior = axis_quantile_mask(s2_source, affine, 2, "inferior", 0.45)
    s2 = union_masks(axis_quantile_mask(s2_inferior, affine, 1, "posterior", 0.70), axis_quantile_mask(s2_inferior, affine, 1, "anterior", 0.30))
    add("area_somatosensorial_secundaria", s2, "aparc+aseg + cuantiles", "Aproximación opercular/parietal inferior derivada de postcentral y supramarginal.")

    lob6_base = cerebellum_cortex & (~vermis)
    lob6 = axis_quantile_mask(lob6_base, affine, 2, "superior", 0.42) & axis_quantile_mask(lob6_base, affine, 1, "anterior", 0.60)
    add("cerebelo_lobulo_vi", lob6, "aseg + cuantiles", "Sector superior-anterior del hemisferio cerebeloso.")

    add("area_motora_presuplementaria", axis_quantile_mask(sma_medial, affine, 1, "anterior", 0.45), "aparc+aseg + cuantiles", "Porción medial anterior del complejo superior frontal/paracentral.")
    add("area_somatosensorial_secundaria_posterior", axis_quantile_mask(s2, affine, 1, "posterior", 0.50), "aparc+aseg + cuantiles", "Mitad posterior de la aproximación S2.")
    add("giro_temporal_inferior_posterior", axis_quantile_mask(inf_temp, affine, 1, "posterior", 0.50), "aparc+aseg + cuantiles", "Mitad posterior del giro temporal inferior bilateral.")
    return regions


###############################################################################################################################
###############################################################################################################################

# ══════════════════════════════════════════════════════════════════════════════
#  7. GUARDADO DE MÁSCARAS, LABELMAPS Y VOLUMETRÍA
# ══════════════════════════════════════════════════════════════════════════════


def save_region_masks(
    *,
    project_root: Path,
    patient: str,
    stage: str,
    subject_id: str,
    t1_path: Path,
    aparc_t1: Path,
    aseg_t1: Path,
    output_dir: Path,
) -> dict:
    output_dir.mkdir(parents=True, exist_ok=True)
    individual_dir = output_dir / "individuales"
    individual_dir.mkdir(parents=True, exist_ok=True)

    t1_img, _ = load_nifti(t1_path)
    aparc_img, aparc_data = load_nifti(aparc_t1)
    aseg_img, aseg_data = load_nifti(aseg_t1)

    if not geometry_matches(t1_img, aparc_img):
        raise RuntimeError(f"aparc_aseg_T1 no coincide con T1: {aparc_t1}")
    if not geometry_matches(t1_img, aseg_img):
        raise RuntimeError(f"aseg_T1 no coincide con T1: {aseg_t1}")

    regions = build_regions(aparc_data.astype(np.int32), aseg_data.astype(np.int32), t1_img)
    vvox = voxel_volume_mm3(t1_img)
    brain_mask = aseg_data > 0
    midline_x = estimate_midline_x(brain_mask, t1_img.affine)

    combined = np.zeros(t1_img.shape[:3], dtype=np.uint16)
    combined_left = np.zeros(t1_img.shape[:3], dtype=np.uint16)
    combined_right = np.zeros(t1_img.shape[:3], dtype=np.uint16)
    rows: list[dict] = []

    for rg in regions:
        mask = rg.mask.astype(bool)
        name = sanitize_name(rg.name)
        bilateral_path = individual_dir / f"{name}.nii"
        save_nifti(t1_img, mask.astype(np.uint8), bilateral_path, dtype=np.uint8)

        total_vox = int(mask.sum())
        total_mm3 = total_vox * vvox
        left_path = right_path = ""
        left_vox = right_vox = ""
        left_mm3 = right_mm3 = ""

        if name not in MIDLINE_REGIONS:
            left_mask = hemispheric_split(mask, t1_img.affine, "left", midline_x, tol=0.0)
            right_mask = hemispheric_split(mask, t1_img.affine, "right", midline_x, tol=0.0)
            left_path_p = individual_dir / f"{name}_izquierda.nii"
            right_path_p = individual_dir / f"{name}_derecha.nii"
            save_nifti(t1_img, left_mask.astype(np.uint8), left_path_p, dtype=np.uint8)
            save_nifti(t1_img, right_mask.astype(np.uint8), right_path_p, dtype=np.uint8)
            left_path = str(left_path_p)
            right_path = str(right_path_p)
            left_vox = int(left_mask.sum())
            right_vox = int(right_mask.sum())
            left_mm3 = float(left_vox * vvox)
            right_mm3 = float(right_vox * vvox)
            combined_left[left_mask & (combined_left == 0)] = rg.label_id
            combined_right[right_mask & (combined_right == 0)] = rg.label_id

        combined[mask & (combined == 0)] = rg.label_id
        rows.append({
            "region_id": rg.label_id,
            "region_name": name,
            "voxels_total": total_vox,
            "volume_total_mm3": round(total_mm3, 3),
            "volume_total_ml": round(total_mm3 / 1000.0, 6),
            "voxels_izquierda": left_vox,
            "volume_izquierda_mm3": round(left_mm3, 3) if isinstance(left_mm3, float) else "",
            "volume_izquierda_ml": round(left_mm3 / 1000.0, 6) if isinstance(left_mm3, float) else "",
            "voxels_derecha": right_vox,
            "volume_derecha_mm3": round(right_mm3, 3) if isinstance(right_mm3, float) else "",
            "volume_derecha_ml": round(right_mm3 / 1000.0, 6) if isinstance(right_mm3, float) else "",
            "source": rg.source,
            "note": rg.note,
            "mask_bilateral": str(bilateral_path),
            "mask_izquierda": left_path,
            "mask_derecha": right_path,
        })

    combined_path = output_dir / "regiones_adicionales_multietiqueta_T1.nii"
    combined_left_path = output_dir / "regiones_adicionales_multietiqueta_izquierda_T1.nii"
    combined_right_path = output_dir / "regiones_adicionales_multietiqueta_derecha_T1.nii"
    labels_path = output_dir / "regiones_adicionales_labels_itksnap.txt"
    csv_path = output_dir / "volumen_regiones_adicionales_por_hemisferio.csv"
    json_path = output_dir / "regiones_adicionales_reporte.json"

    save_nifti(t1_img, combined, combined_path, dtype=np.uint16)
    save_nifti(t1_img, combined_left, combined_left_path, dtype=np.uint16)
    save_nifti(t1_img, combined_right, combined_right_path, dtype=np.uint16)
    write_itksnap_labels(labels_path, regions)

    with csv_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)

    report = {
        "project_root": str(project_root),
        "patient": patient,
        "timepoint": stage,
        "subject_id": subject_id,
        "reference_t1": str(t1_path),
        "sources": {"aparc_aseg_t1": str(aparc_t1), "aseg_t1": str(aseg_t1)},
        "outputs": {
            "combined_labelmap": str(combined_path),
            "combined_labelmap_left": str(combined_left_path),
            "combined_labelmap_right": str(combined_right_path),
            "label_descriptions": str(labels_path),
            "volume_csv": str(csv_path),
            "individual_dir": str(individual_dir),
        },
        "regions": rows,
    }
    json_path.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")

    (output_dir / "LEEME_REGIONES_ADICIONALES.txt").write_text(
        "REGIONES ADICIONALES EN ESPACIO T1\n\n"
        f"Paciente: {patient}\nEtapa: {stage}\nT1: {t1_path}\n\n"
        f"Labelmap bilateral: {combined_path}\n"
        f"Tabla de etiquetas: {labels_path}\n"
        f"Máscaras individuales: {individual_dir}\n"
        f"CSV volúmenes: {csv_path}\n\n"
        "En ITK-SNAP: abrir T1 como Main Image, abrir el labelmap como Segmentation y cargar la tabla de etiquetas.\n",
        encoding="utf-8",
    )
    return report


def generate_19_zones_for_case(cfg: PipelineConfig, patient: str, stage: str) -> Optional[dict]:
    subject_dir = cfg.data_root() / patient
    stage_dir = resolve_stage_dir(subject_dir, stage)
    t1_path = first_existing(t1_candidates(stage_dir))
    out_dir = cfg.results_root() / patient / stage / "tractografia_propia" / "regiones_adicionales_T1"
    if t1_path is None:
        log(f"[ZONAS] {patient}/{stage}: no se encontró rT1/rAnatomico. Se omite.")
        return None
    log("\n" + "=" * 92)
    log(f"[{stamp()}] 19 ZONAS · {patient} / {stage}")
    log("=" * 92)
    log(f"T1: {t1_path}")
    log(f"Salida: {out_dir}")
    sid = subject_id_for(patient, stage)
    aparc_t1, aseg_t1 = ensure_freesurfer_labels_in_t1(cfg=cfg, patient=patient, stage=stage, t1_path=t1_path, output_dir=out_dir)
    return save_region_masks(
        project_root=cfg.project_root,
        patient=patient,
        stage=stage,
        subject_id=sid,
        t1_path=t1_path,
        aparc_t1=aparc_t1,
        aseg_t1=aseg_t1,
        output_dir=out_dir,
    )


###############################################################################################################################
###############################################################################################################################

# ══════════════════════════════════════════════════════════════════════════════
#  8. REGISTRO, NORMALIZACIÓN Y CORRELACIÓN LOCAL ANTES/DESPUÉS
# ══════════════════════════════════════════════════════════════════════════════


def register_after_to_before(before_path: Path, after_path: Path, output_path: Path) -> Path:
    if output_path.exists():
        log(f"[CORRELACIÓN] Reutilizando registro: {output_path}")
        return output_path
    log(f"[{stamp()}] Registrando Después -> Antes")
    fixed = sitk.ReadImage(str(before_path), sitk.sitkFloat32)
    moving = sitk.ReadImage(str(after_path), sitk.sitkFloat32)
    initial = sitk.CenteredTransformInitializer(
        fixed,
        moving,
        sitk.Euler3DTransform(),
        sitk.CenteredTransformInitializerFilter.GEOMETRY,
    )
    reg = sitk.ImageRegistrationMethod()
    reg.SetMetricAsMattesMutualInformation(numberOfHistogramBins=50)
    reg.SetMetricSamplingStrategy(reg.RANDOM)
    reg.SetMetricSamplingPercentage(0.20, seed=42)
    reg.SetInterpolator(sitk.sitkLinear)
    reg.SetOptimizerAsGradientDescent(
        learningRate=1.0,
        numberOfIterations=250,
        convergenceMinimumValue=1e-6,
        convergenceWindowSize=15,
    )
    reg.SetOptimizerScalesFromPhysicalShift()
    reg.SetShrinkFactorsPerLevel([4, 2, 1])
    reg.SetSmoothingSigmasPerLevel([2, 1, 0])
    reg.SmoothingSigmasAreSpecifiedInPhysicalUnitsOn()
    reg.SetInitialTransform(initial, inPlace=False)
    transform = reg.Execute(fixed, moving)
    registered = sitk.Resample(moving, fixed, transform, sitk.sitkLinear, 0.0, sitk.sitkFloat32)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    sitk.WriteImage(registered, str(output_path))
    sitk.WriteTransform(transform, str(output_path.parent / "registro_despues_a_antes.tfm"))
    log(f"[CORRELACIÓN] Registro terminado. Métrica final: {reg.GetMetricValue():.6f}")
    return output_path


def robust_zscore(data: np.ndarray, mask: np.ndarray) -> np.ndarray:
    vals = data[mask]
    if vals.size == 0:
        raise RuntimeError("La máscara de normalización está vacía.")
    med = np.median(vals)
    mad = np.median(np.abs(vals - med))
    scale = 1.4826 * mad
    if scale < 1e-8:
        scale = np.std(vals)
    if scale < 1e-8:
        scale = 1.0
    return ((data - med) / scale).astype(np.float32)


def local_corr_2d_by_axis(a: np.ndarray, b: np.ndarray, axis: int, window: int) -> tuple[np.ndarray, np.ndarray]:
    if window % 2 == 0:
        raise ValueError("WINDOW_SIZE debe ser impar.")
    if axis == 0:
        size = (1, window, window)
    elif axis == 1:
        size = (window, 1, window)
    elif axis == 2:
        size = (window, window, 1)
    else:
        raise ValueError("axis debe ser 0, 1 o 2")
    a = a.astype(np.float32, copy=False)
    b = b.astype(np.float32, copy=False)
    ma = uniform_filter(a, size=size, mode="reflect")
    mb = uniform_filter(b, size=size, mode="reflect")
    maa = uniform_filter(a * a, size=size, mode="reflect")
    mbb = uniform_filter(b * b, size=size, mode="reflect")
    mab = uniform_filter(a * b, size=size, mode="reflect")
    va = np.maximum(maa - ma * ma, 0.0)
    vb = np.maximum(mbb - mb * mb, 0.0)
    cov = mab - ma * mb
    den = np.sqrt(va * vb)
    valid = den > 1e-6
    corr = np.zeros_like(a, dtype=np.float32)
    corr[valid] = cov[valid] / den[valid]
    corr = np.clip(corr, -1.0, 1.0)
    return corr, valid


def local_corr_3_planes(a: np.ndarray, b: np.ndarray, mask: np.ndarray, window: int) -> tuple[np.ndarray, np.ndarray, Dict[str, np.ndarray]]:
    plane_maps: Dict[str, np.ndarray] = {}
    valid_maps: Dict[str, np.ndarray] = {}
    for axis, name in [(0, "eje_0"), (1, "eje_1"), (2, "eje_2")]:
        log(f"[CORRELACIÓN] Ventanas 2D en todos los cortes del {name}...")
        corr, valid = local_corr_2d_by_axis(a, b, axis, window)
        valid &= mask
        corr[~valid] = 0.0
        plane_maps[name] = corr
        valid_maps[name] = valid
    suma = np.zeros_like(a, dtype=np.float32)
    cuenta = np.zeros_like(a, dtype=np.float32)
    for name, mapa in plane_maps.items():
        v = valid_maps[name]
        suma[v] += mapa[v]
        cuenta[v] += 1.0
    valid_final = cuenta > 0
    combined = np.zeros_like(a, dtype=np.float32)
    combined[valid_final] = suma[valid_final] / cuenta[valid_final]
    combined[~mask] = 0.0
    return combined, valid_final, plane_maps


def safe_pearson(a: np.ndarray, b: np.ndarray) -> float:
    if a.size < 3 or np.std(a) < 1e-8 or np.std(b) < 1e-8:
        return float("nan")
    return float(pearsonr(a, b).statistic)


def region_metrics(before_z: np.ndarray, after_z: np.ndarray, corr_map: np.ndarray, corr_valid: np.ndarray, mask: np.ndarray) -> dict:
    valid = mask & corr_valid & np.isfinite(before_z) & np.isfinite(after_z) & np.isfinite(corr_map)
    n = int(valid.sum())
    if n < 3:
        return {
            "voxels": n,
            "mean_before_z": float("nan"),
            "mean_after_z": float("nan"),
            "pearson_r": float("nan"),
            "local_corr_mean": float("nan"),
            "local_corr_median": float("nan"),
            "local_corr_p10": float("nan"),
            "low_corr_percent": float("nan"),
            "mae_z": float("nan"),
            "rmse_z": float("nan"),
        }
    a = before_z[valid]
    b = after_z[valid]
    c = corr_map[valid]
    diff = b - a
    return {
        "voxels": n,
        "mean_before_z": float(np.mean(a)),
        "mean_after_z": float(np.mean(b)),
        "pearson_r": safe_pearson(a, b),
        "local_corr_mean": float(np.mean(c)),
        "local_corr_median": float(np.median(c)),
        "local_corr_p10": float(np.percentile(c, 10)),
        "low_corr_percent": float(100.0 * np.mean(c < LOW_CORR_THRESHOLD)),
        "mae_z": float(np.mean(np.abs(diff))),
        "rmse_z": float(np.sqrt(np.mean(diff ** 2))),
    }


def find_region_mask(roi_root: Path, region: str) -> Optional[Path]:
    candidates = [
        roi_root / "individuales" / f"{region}.nii",
        roi_root / "individuales" / f"{region}.nii.gz",
        roi_root / f"{region}.nii",
        roi_root / f"{region}.nii.gz",
    ]
    for p in candidates:
        if p.exists():
            return p
    if roi_root.exists():
        found = list(roi_root.rglob(f"{region}.nii")) + list(roi_root.rglob(f"{region}.nii.gz"))
        found = [p for p in found if "_izquierda" not in p.name and "_derecha" not in p.name]
        if found:
            found.sort(key=lambda p: p.stat().st_mtime, reverse=True)
            return found[0]
    return None


def compute_before_after_correlation(cfg: PipelineConfig, patient: str) -> Optional[dict]:
    before_stage = next((s for s in cfg.stages if stage_key(s) == "antes"), "Antes")
    after_stage = next((s for s in cfg.stages if stage_key(s) == "despues"), "Despues")
    if not any(stage_key(s) == "antes" for s in cfg.stages) or not any(stage_key(s) == "despues" for s in cfg.stages):
        log(f"[CORRELACIÓN] {patient}: se requieren etapas Antes y Despues. Usa --stages Antes Despues.")
        return None

    before_dir = resolve_stage_dir(cfg.data_root() / patient, before_stage)
    after_dir = resolve_stage_dir(cfg.data_root() / patient, after_stage)
    before_path = first_existing(t1_candidates(before_dir))
    after_path = first_existing(t1_candidates(after_dir))
    if before_path is None or after_path is None:
        log(f"[CORRELACIÓN] {patient}: faltan imágenes Antes/Después. Se omite.")
        return None

    roi_root = cfg.results_root() / patient / before_stage / "tractografia_propia" / "regiones_adicionales_T1"
    out_dir = cfg.results_root() / patient / "correlacion_antes_despues_19_zonas"
    out_dir.mkdir(parents=True, exist_ok=True)

    log("\n" + "=" * 92)
    log(f"[{stamp()}] CORRELACIÓN 19 ZONAS · {patient}")
    log("=" * 92)
    log(f"ANTES:   {before_path}")
    log(f"DESPUÉS: {after_path}")
    log(f"ROIs:    {roi_root}")
    log(f"Salida:  {out_dir}")

    after_registered = register_after_to_before(before_path, after_path, out_dir / "despues_registrado_a_antes.nii.gz")
    before_img, before = load_nifti(before_path)
    after_img, after = load_nifti(after_registered)
    before = before.astype(np.float32)
    after = after.astype(np.float32)
    if not geometry_matches(before_img, after_img):
        raise RuntimeError("Después registrado no coincide geométricamente con Antes.")

    region_masks: Dict[str, np.ndarray] = {}
    analysis_mask = np.zeros(before.shape, dtype=bool)
    for region, _, _ in REGION_ORDER:
        mask_path = find_region_mask(roi_root, region)
        if mask_path is None:
            log(f"[ADVERTENCIA] Falta máscara: {region}")
            continue
        mask_img, mask_data = load_nifti(mask_path)
        if not geometry_matches(before_img, mask_img):
            raise RuntimeError(f"Máscara incompatible con Antes: {mask_path}")
        m = mask_data > 0
        region_masks[region] = m
        analysis_mask |= m

    if int(analysis_mask.sum()) == 0:
        raise RuntimeError("No se encontró ninguna de las 19 máscaras para correlación.")

    before_z = robust_zscore(before, analysis_mask)
    after_z = robust_zscore(after, analysis_mask)
    corr_map, corr_valid, plane_maps = local_corr_3_planes(before_z, after_z, analysis_mask, WINDOW_SIZE)
    low_mask = analysis_mask & corr_valid & (corr_map < LOW_CORR_THRESHOLD)
    low_score = np.zeros_like(corr_map, dtype=np.float32)
    low_score[corr_valid] = 1.0 - corr_map[corr_valid]

    low_labelmap = np.zeros(before.shape, dtype=np.uint16)
    for region, label_id, _ in REGION_ORDER:
        m = region_masks.get(region)
        if m is None:
            continue
        low_labelmap[m & low_mask & (low_labelmap == 0)] = label_id

    save_nifti(before_img, corr_map, out_dir / "correlacion_local_3planos_T1.nii.gz", dtype=np.float32)
    save_nifti(before_img, low_score, out_dir / "baja_correlacion_score_T1.nii.gz", dtype=np.float32)
    save_nifti(before_img, low_mask.astype(np.uint8), out_dir / "baja_correlacion_mask_T1.nii.gz", dtype=np.uint8)
    save_nifti(before_img, low_labelmap, out_dir / "baja_correlacion_zonas_labelmap_T1.nii.gz", dtype=np.uint16)
    for name, mapa in plane_maps.items():
        save_nifti(before_img, mapa, out_dir / f"correlacion_local_{name}_T1.nii.gz", dtype=np.float32)
    write_itksnap_labels(out_dir / "baja_correlacion_zonas_labels_itksnap.txt")

    metrics: Dict[str, Optional[dict]] = {}
    for region, _, _ in REGION_ORDER:
        m = region_masks.get(region)
        metrics[region] = None if m is None else region_metrics(before_z, after_z, corr_map, corr_valid, m)

    report = {
        "patient": patient,
        "before": str(before_path),
        "after": str(after_path),
        "after_registered": str(after_registered),
        "reference_space": "Antes",
        "window_size_2d": WINDOW_SIZE,
        "low_corr_threshold": LOW_CORR_THRESHOLD,
        "method": "Rigid registration + robust z-score + local 2D Pearson/NCC in three orthogonal planes",
        "outputs": {
            "correlation_map": str(out_dir / "correlacion_local_3planos_T1.nii.gz"),
            "low_correlation_score": str(out_dir / "baja_correlacion_score_T1.nii.gz"),
            "low_correlation_labelmap": str(out_dir / "baja_correlacion_zonas_labelmap_T1.nii.gz"),
            "label_descriptions": str(out_dir / "baja_correlacion_zonas_labels_itksnap.txt"),
        },
        "regions": metrics,
    }
    (out_dir / "reporte_correlacion_antes_despues.json").write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    return report


###############################################################################################################################
###############################################################################################################################

# ══════════════════════════════════════════════════════════════════════════════
#  9. EXCEL DE VOLÚMENES Y CORRELACIONES
# ══════════════════════════════════════════════════════════════════════════════


def excel_value(v):
    if v is None:
        return "NO ENCONTRADO"
    if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
        return "NO CALCULABLE"
    return v


def style_header(ws, row: int, max_col: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    white = "FFFFFF"
    thin = Side(style="thin", color="C9C9C9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col in range(1, max_col + 1):
        c = ws.cell(row=row, column=col)
        c.fill = fill
        c.font = Font(color=white, bold=True, size=11)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border


def save_workbook(wb: Workbook, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    log(f"[EXCEL] Guardado: {path}")
    return path


def build_volume_excel(cfg: PipelineConfig) -> Optional[Path]:
    wb = Workbook()
    ws = wb.active
    ws.title = "Volúmenes 19 zonas"
    headers = [
        "Paciente", "Etapa", "ID", "Zona", "Voxeles total", "Volumen total mm3", "Volumen total ml",
        "Voxeles izquierda", "Volumen izquierda mm3", "Volumen izquierda ml",
        "Voxeles derecha", "Volumen derecha mm3", "Volumen derecha ml", "Fuente", "Nota",
    ]
    ws.append(headers)
    style_header(ws, 1, len(headers))
    for patient in cfg.patients:
        for stage in cfg.stages:
            report_path = cfg.results_root() / patient / stage / "tractografia_propia" / "regiones_adicionales_T1" / "regiones_adicionales_reporte.json"
            if not report_path.exists():
                continue
            data = json.loads(report_path.read_text(encoding="utf-8"))
            for rg in data.get("regions", []):
                ws.append([
                    patient, stage, rg.get("region_id"), rg.get("region_name"), rg.get("voxels_total"),
                    rg.get("volume_total_mm3"), rg.get("volume_total_ml"), rg.get("voxels_izquierda"),
                    rg.get("volume_izquierda_mm3"), rg.get("volume_izquierda_ml"), rg.get("voxels_derecha"),
                    rg.get("volume_derecha_mm3"), rg.get("volume_derecha_ml"), rg.get("source"), rg.get("note"),
                ])
    if ws.max_row == 1:
        return None
    widths = {"A": 16, "B": 12, "C": 8, "D": 42, "E": 15, "F": 18, "G": 16, "H": 18, "I": 20, "J": 18, "K": 18, "L": 20, "M": 18, "N": 25, "O": 70}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    return save_workbook(wb, cfg.results_root() / "resumen_volumenes_19_zonas.xlsx")


def build_correlation_excel(cfg: PipelineConfig, reports: list[dict]) -> Optional[Path]:
    if not reports:
        return None
    wb = Workbook()
    ws = wb.active
    ws.title = "Correlación por zona"
    headers = [
        "Paciente", "Zona", "Antes media z", "Después media z", "Pearson r", "Corr local media",
        "Corr local mediana", "P10 corr local", "% baja correlación", "MAE z", "RMSE z", "Voxeles",
    ]
    ws.append(headers)
    style_header(ws, 1, len(headers))
    low_fill = PatternFill("solid", fgColor="F4CCCC")
    warn_fill = PatternFill("solid", fgColor="FFF2CC")
    thin = Side(style="thin", color="C9C9C9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for report in reports:
        patient = report.get("patient")
        metrics = report.get("regions", {})
        for region, _, _ in REGION_ORDER:
            m = metrics.get(region)
            row = [
                patient, region,
                None if m is None else m.get("mean_before_z"),
                None if m is None else m.get("mean_after_z"),
                None if m is None else m.get("pearson_r"),
                None if m is None else m.get("local_corr_mean"),
                None if m is None else m.get("local_corr_median"),
                None if m is None else m.get("local_corr_p10"),
                None if m is None else m.get("low_corr_percent"),
                None if m is None else m.get("mae_z"),
                None if m is None else m.get("rmse_z"),
                None if m is None else m.get("voxels"),
            ]
            ws.append([excel_value(v) for v in row])
    for rows in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for c in rows:
            c.border = border
            c.alignment = Alignment(vertical="center", wrap_text=True)
            if isinstance(c.value, str) and c.value in {"NO ENCONTRADO", "NO CALCULABLE"}:
                c.fill = warn_fill
        c_corr = rows[5]
        if isinstance(c_corr.value, (int, float)) and c_corr.value < LOW_CORR_THRESHOLD:
            c_corr.fill = low_fill
    for col, width in {"A": 16, "B": 42, "C": 18, "D": 18, "E": 13, "F": 18, "G": 20, "H": 18, "I": 18, "J": 13, "K": 13, "L": 12}.items():
        ws.column_dimensions[col].width = width
    return save_workbook(wb, cfg.results_root() / "resumen_correlacion_19_zonas.xlsx")


###############################################################################################################################
###############################################################################################################################

# ══════════════════════════════════════════════════════════════════════════════
#  10. FUNCIÓN PÚBLICA PARA MAIN.PY
# ══════════════════════════════════════════════════════════════════════════════


def run_zonas_correlacion_module(cfg: PipelineConfig) -> pd.DataFrame:
    """Ejecuta delimitación de 19 zonas, correlación Antes/Después y Excel.

    Esta función está diseñada para integrarse con `python main.py run --sections
    zonas_correlacion`. La delimitación de zonas corre por paciente y etapa. La
    correlación longitudinal se calcula cuando la corrida incluye Antes y Después.
    """
    rows: list[dict] = []
    correlation_reports: list[dict] = []

    for patient in cfg.patients:
        for stage in cfg.stages:
            row = {"patient": patient, "stage": stage, "section": "zonas_correlacion", "status": "pendiente"}
            try:
                report = generate_19_zones_for_case(cfg, patient, stage)
                if report is None:
                    row["status"] = "omitido_sin_t1"
                else:
                    row["status"] = "zonas_ok"
                    row["regions_report"] = str(cfg.results_root() / patient / stage / "tractografia_propia" / "regiones_adicionales_T1" / "regiones_adicionales_reporte.json")
            except Exception as exc:
                row["status"] = "fallo_zonas"
                row["error"] = str(exc)
                log(f"[ERROR] 19 zonas {patient}/{stage}: {exc}")
            rows.append(row)

        try:
            corr_report = compute_before_after_correlation(cfg, patient)
            if corr_report is not None:
                correlation_reports.append(corr_report)
                rows.append({"patient": patient, "stage": "Antes_vs_Despues", "section": "zonas_correlacion", "status": "correlacion_ok"})
        except Exception as exc:
            rows.append({"patient": patient, "stage": "Antes_vs_Despues", "section": "zonas_correlacion", "status": "fallo_correlacion", "error": str(exc)})
            log(f"[ERROR] Correlación 19 zonas {patient}: {exc}")

    volume_excel = build_volume_excel(cfg)
    corr_excel = build_correlation_excel(cfg, correlation_reports)
    if volume_excel:
        rows.append({"patient": "todos", "stage": "todas", "section": "zonas_correlacion", "status": "excel_volumenes_ok", "path": str(volume_excel)})
    if corr_excel:
        rows.append({"patient": "todos", "stage": "Antes_vs_Despues", "section": "zonas_correlacion", "status": "excel_correlacion_ok", "path": str(corr_excel)})

    df = pd.DataFrame(rows)
    out_csv = cfg.results_root() / "resumen_modulo_19_zonas_correlacion.csv"
    out_csv.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(out_csv, index=False, encoding="utf-8-sig")
    log(f"[RESUMEN] {out_csv}")
    return df
